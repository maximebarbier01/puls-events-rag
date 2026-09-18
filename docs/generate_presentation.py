"""
Génère la présentation PowerPoint du projet (docs/Puls-Events-RAG-Presentation.pptx).

Outil ponctuel de génération de contenu, pas une dépendance du système RAG lui-même :
python-pptx n'est donc pas dans pyproject.toml, à installer à part si besoin :
    pip install python-pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Cm, Emu, Pt

# --- Palette sobre et professionnelle -----------------------------------------
DARK_BLUE = RGBColor(0x1B, 0x2A, 0x4A)
ACCENT = RGBColor(0x2E, 0x7D, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF2, 0xF3, 0xF5)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MEDIUM_GRAY = RGBColor(0x6B, 0x6B, 0x6B)

SLIDE_W = Cm(33.87)  # 16:9
SLIDE_H = Cm(19.05)

OUTPUT_PATH = Path(__file__).resolve().parent / "Puls-Events-RAG-Presentation.pptx"


def blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])  # layout vide


def set_background(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    return box


def add_bullets(slide, left, top, width, height, items, size=20, color=DARK_GRAY):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(12)
    return box


def add_title_bar(slide, title, kicker=None):
    set_background(slide, WHITE)
    add_textbox(slide, Cm(1.2), Cm(0.8), Cm(30), Cm(1.6), title, 30, DARK_BLUE, bold=True)
    if kicker:
        add_textbox(slide, Cm(1.2), Cm(2.2), Cm(30), Cm(0.8), kicker, 14, ACCENT, bold=True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(1.2), Cm(2.9), Cm(6), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def content_slide(prs, title, bullets, kicker=None, footer=None):
    slide = blank_slide(prs)
    add_title_bar(slide, title, kicker)
    add_bullets(slide, Cm(1.4), Cm(3.6), Cm(30), Cm(13), bullets)
    if footer:
        add_textbox(slide, Cm(1.2), Cm(17.8), Cm(30), Cm(1), footer, 12, MEDIUM_GRAY)
    return slide


def section_slide(prs, title, subtitle=None):
    slide = blank_slide(prs)
    set_background(slide, DARK_BLUE)
    add_textbox(slide, Cm(2), Cm(8), Cm(29.87), Cm(3), title, 40, WHITE, bold=True, align=PP_ALIGN.CENTER)
    if subtitle:
        add_textbox(slide, Cm(2), Cm(11), Cm(29.87), Cm(1.5), subtitle, 18, RGBColor(0xC9, 0xD3, 0xE0), align=PP_ALIGN.CENTER)
    return slide


def big_number_box(slide, left, top, width, number, label, number_color=ACCENT):
    add_textbox(slide, left, top, width, Cm(2.5), number, 44, number_color, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, left, top + Cm(2.3), width, Cm(1.5), label, 14, DARK_GRAY, align=PP_ALIGN.CENTER)


# --------------------------------------------------------------------------- #

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

# 1. Titre --------------------------------------------------------------------
slide = blank_slide(prs)
set_background(slide, DARK_BLUE)
add_textbox(slide, Cm(2), Cm(6.5), Cm(29.87), Cm(2.5), "Puls-Events RAG", 44, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_textbox(slide, Cm(2), Cm(8.8), Cm(29.87), Cm(1.5), "Assistant intelligent de recommandation d'événements culturels", 20, RGBColor(0xC9, 0xD3, 0xE0), align=PP_ALIGN.CENTER)
add_textbox(slide, Cm(2), Cm(10.6), Cm(29.87), Cm(1), "POC — LangChain · FAISS · Mistral", 16, ACCENT, align=PP_ALIGN.CENTER)
add_textbox(slide, Cm(2), Cm(16.5), Cm(29.87), Cm(1), "Maxime Barbier — Parcours OpenClassrooms « Concevez et déployez un système RAG »", 13, RGBColor(0x9A, 0xA7, 0xBD), align=PP_ALIGN.CENTER)

# 2. Contexte & problématique ---------------------------------------------------
content_slide(
    prs,
    "Contexte & problématique",
    [
        "Puls-Events : plateforme de recommandations culturelles personnalisées",
        "Besoin métier : répondre en langage naturel aux questions des utilisateurs sur les événements à venir",
        "Problématique : comment un système RAG peut-il combiner recherche sémantique et génération de réponse pour rendre cette recommandation naturelle et fiable ?",
        "Mission confiée par Jérémy, responsable technique Puls-Events",
    ],
    kicker="LE PROBLÈME",
)

# 3. Objectifs du POC -----------------------------------------------------------
content_slide(
    prs,
    "Objectifs du POC",
    [
        "Démontrer la faisabilité technique : LangChain + FAISS + Mistral, bout en bout",
        "Démontrer la valeur métier : réponses pertinentes, fiables, sans hallucination",
        "Démontrer la performance : API testable en local, résultats mesurés objectivement (Ragas)",
        "Périmètre : événements culturels en Moselle, sur une fenêtre d'un an d'historique + tout l'avenir",
    ],
    kicker="LE PROBLÈME",
)

# 4. Architecture globale ---------------------------------------------------------
slide = blank_slide(prs)
add_title_bar(slide, "Architecture globale", kicker="LA SOLUTION")

boxes_data = [
    "Open Agenda\n(API)",
    "Nettoyage &\nfiltrage",
    "Chunking &\nembeddings",
    "Index\nFAISS",
    "LangChain\n(retrieval)",
    "Mistral\n(génération)",
    "API REST\n(FastAPI)",
]
n = len(boxes_data)
box_w = Cm(3.9)
box_h = Cm(2.6)
gap = Cm(0.55)
total_w = n * box_w + (n - 1) * gap
start_x = (SLIDE_W - total_w) // 2
y = Cm(8.5)

shapes = []
for i, label in enumerate(boxes_data):
    x = start_x + i * (box_w + gap)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_w, box_h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_BLUE if i not in (3,) else ACCENT
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = label
    p.font.size = Pt(13)
    p.font.color.rgb = WHITE
    p.font.bold = True
    shapes.append(shape)

for i in range(n - 1):
    x1 = start_x + i * (box_w + gap) + box_w
    x2 = x1 + gap
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y + box_h // 2, x2, y + box_h // 2)
    connector.line.color.rgb = MEDIUM_GRAY
    connector.line.width = Pt(2)

add_bullets(
    slide, Cm(1.4), Cm(12.2), Cm(30), Cm(5),
    [
        "Pipeline scripté de bout en bout (scripts/00 à 04), reproductible à la demande",
        "Séparation stricte logique métier (app/) / points d'entrée (scripts/, API)",
        "Endpoint /rebuild : ré-indexation à la demande sans redémarrer le serveur",
    ],
    size=16,
)

# 5. Périmètre des données --------------------------------------------------------
slide = blank_slide(prs)
add_title_bar(slide, "Périmètre des données", kicker="LA SOLUTION")

funnel = [("9083", "événements Moselle\n(sans filtre)"), ("2829", "après filtre géo +\nfenêtre 1 an"), ("1483", "événements culturels\npropres, indexés")]
fw = Cm(9)
fx0 = (SLIDE_W - (fw * 3 + Cm(1) * 2)) // 2
for i, (num, label) in enumerate(funnel):
    x = fx0 + i * (fw + Cm(1))
    big_number_box(slide, x, Cm(5), fw, num, label, number_color=ACCENT if i < 2 else DARK_BLUE)
    if i < 2:
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + fw, Cm(5.6), Cm(0.9), Cm(1))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = MEDIUM_GRAY
        arrow.line.fill.background()

add_bullets(
    slide, Cm(1.4), Cm(10.2), Cm(30), Cm(7),
    [
        "Zone : département de la Moselle (Grand Est), via l'API Explore d'OpenDataSoft",
        "Fenêtre temporelle : 1 an d'historique + tous les événements à venir",
        "Filtre thématique : exclusion explicite des sources non-culturelles (ex : sessions de recrutement France Travail, ~35% du volume brut)",
        "Statut : événements annulés exclus, événements complets gardés mais signalés",
    ],
    size=17,
)

# 6. Pipeline de données --------------------------------------------------------
content_slide(
    prs,
    "Pipeline de données",
    [
        "00 — Récupération via l'API Explore d'OpenDataSoft (paginée, filtrée par département + date)",
        "01 — Nettoyage : suppression HTML, filtre statut/thématique/récence, construction du texte source",
        "02 — Chunking (1000 car. / 150 de chevauchement) + vectorisation + indexation FAISS",
        "Entièrement scripté et reproductible — aucune étape manuelle une fois la clé API configurée",
    ],
    kicker="LA SOLUTION",
)

# 7. Vectorisation & FAISS --------------------------------------------------------
content_slide(
    prs,
    "Vectorisation & FAISS",
    [
        "Embeddings : sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2 (local, HuggingFace)",
        "Choix motivé : gratuit, pas de clé API nécessaire pour indexer/reconstruire, adapté au français, reproductible sans quota",
        "1483 événements → 1920 chunks indexés (index FAISS plat, exact, adapté à ce volume)",
        "Métadonnées conservées par chunk : dates, lieu, coordonnées, URL source, statut complet",
    ],
    kicker="LA SOLUTION",
)

# 8. La chaîne RAG ---------------------------------------------------------------
content_slide(
    prs,
    "La chaîne RAG",
    [
        "Retrieval : top-k (k=5) événements les plus proches sémantiquement de la question",
        "Prompt système strict : répondre uniquement à partir du contexte fourni, refuser explicitement si rien ne correspond",
        "Génération : Mistral (mistral-small-latest) — rapide, peu coûteux, suffisant pour reformuler des infos d'événements",
        "Pas d'historique de conversation (hors périmètre du POC, conformément au sujet)",
    ],
    kicker="LA SOLUTION",
)

# 9. L'API REST -------------------------------------------------------------------
content_slide(
    prs,
    "L'API REST",
    [
        "FastAPI — documentation Swagger interactive générée automatiquement (/docs)",
        "POST /ask — pose une question, reçoit une réponse augmentée + les sources utilisées",
        "POST /rebuild — reconstruit l'index à la demande, protégé par un jeton (header X-Rebuild-Token)",
        "Modèle et index chargés une seule fois au démarrage (pas de rechargement à chaque appel)",
        "Logique métier (app/) totalement découplée du code API (app/api/)",
    ],
    kicker="LA SOLUTION",
)

# 10. Démo live ---------------------------------------------------------------
section_slide(prs, "Démonstration live", "Question posée → réponse générée, en direct")

# 11. Évaluation — méthodologie -----------------------------------------------
content_slide(
    prs,
    "Évaluation — méthodologie",
    [
        "Jeu de test annoté : 12 questions avec réponses de référence (eval/qa_dataset.json)",
        "Inclut des cas pièges volontaires : questions hors périmètre géographique/thématique, pour vérifier le refus d'halluciner",
        "4 métriques Ragas : Faithfulness, Answer Relevancy, Context Precision, Context Recall",
        "Jugement assuré par notre propre modèle Mistral (pas OpenAI) — cohérence de la stack technique",
    ],
    kicker="LES RÉSULTATS",
)

# 12. Évaluation — résultats ----------------------------------------------------
slide = blank_slide(prs)
add_title_bar(slide, "Évaluation — résultats", kicker="LES RÉSULTATS")

table_rows = [
    ("Métrique", "Score global", "Hors questions pièges"),
    ("Faithfulness", "0.888", "0.895"),
    ("Answer relevancy", "0.548", "0.822"),
    ("Context precision", "0.660", "0.667"),
    ("Context recall", "0.872", "0.933"),
]
rows, cols = len(table_rows), 3
table_shape = slide.shapes.add_table(rows, cols, Cm(3), Cm(4), Cm(28), Cm(8))
table = table_shape.table
table.columns[0].width = Cm(12)
table.columns[1].width = Cm(8)
table.columns[2].width = Cm(8)
for r, row in enumerate(table_rows):
    for c, val in enumerate(row):
        cell = table.cell(r, c)
        cell.text = val
        para = cell.text_frame.paragraphs[0]
        para.font.size = Pt(16)
        para.font.bold = r == 0
        para.font.color.rgb = WHITE if r == 0 else DARK_GRAY
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE if r == 0 else (LIGHT_GRAY if r % 2 == 0 else WHITE)

add_bullets(
    slide, Cm(1.4), Cm(13), Cm(30), Cm(5.5),
    [
        "Answer relevancy chute sur les questions pièges : le système refuse correctement, mais la métrique compare par similarité sémantique une réponse de refus courte à la question — limite connue de la métrique, pas un défaut du système",
        "Point faible identifié : context precision (~0.66) — piste d'amélioration à l'étape suivante (ajustement de k, reranking)",
    ],
    size=15,
)

# 13. Conteneurisation Docker -------------------------------------------------
content_slide(
    prs,
    "Conteneurisation Docker",
    [
        "Image légère : build multi-stage Poetry, utilisateur non-root, code + dépendances uniquement",
        "Données et index montés en volumes (jamais figés dans l'image) — toujours à jour, /rebuild fonctionnel",
        "Secrets jamais copiés dans l'image (injectés au runtime via .env)",
        "Une commande pour démarrer : docker compose up --build",
        "CI GitHub Actions : tests automatiques sur chaque Pull Request, évaluation Ragas sur chaque merge vers main",
    ],
    kicker="PRÊT POUR ALLER PLUS LOIN",
)

# 14. Limites & perspectives ----------------------------------------------------
content_slide(
    prs,
    "Limites & perspectives d'amélioration",
    [
        "Pas de filtrage temporel réel : \"ce week-end\"/\"demain\" reposent sur la similarité sémantique, pas sur une vraie date — à corriger par extraction de plage de dates + filtre sur les métadonnées",
        "Context precision perfectible : tester différentes valeurs de k, envisager un reranking",
        "Volumétrie limitée à la Moselle pour le POC — extension géographique possible sans changement d'architecture",
        "Pas de déploiement cloud à ce stade (hors périmètre du POC) — Docker prouve que le système est prêt pour un déploiement élargi",
    ],
    kicker="PRÊT POUR ALLER PLUS LOIN",
)

# 15. Conclusion ------------------------------------------------------------------
slide = blank_slide(prs)
set_background(slide, DARK_BLUE)
add_textbox(slide, Cm(2), Cm(6), Cm(29.87), Cm(2), "Conclusion", 36, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_bullets(
    slide, Cm(6), Cm(9), Cm(21.87), Cm(6),
    [
        "POC fonctionnel et testé de bout en bout : données → index → RAG → API → Docker",
        "Qualité mesurée objectivement (Ragas), pas juste affirmée",
        "Base solide pour un déploiement élargi",
    ],
    size=18,
    color=WHITE,
)
add_textbox(slide, Cm(2), Cm(16.5), Cm(29.87), Cm(1.5), "Questions ?", 24, ACCENT, bold=True, align=PP_ALIGN.CENTER)

prs.save(str(OUTPUT_PATH))
print(f"Présentation générée : {OUTPUT_PATH} ({len(prs.slides)} slides)")
